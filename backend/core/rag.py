"""Org-context retrieval: user documents -> Pinecone -> agent prompt.

This is the qualitative half of the diagnostic. The benchmark side knows what
peers spend; it cannot know that the company runs three shared-service sites or
that a facilities contract has a 90-day exit. Those facts live in documents the
user uploads here, and are retrieved per turn so the agent can ground a lever in
something specific to the business.

Embedding is Pinecone's own (integrated inference): records are upserted as raw
text on a `text` field and the index embeds them server-side. That keeps the
embedding model and its dimension out of this codebase entirely -- no separate
embedding deployment to provision and no dimension to keep in sync.

One namespace per company keeps tenants apart: a search for company A can never
reach company B's documents, because the namespace is the boundary.
"""

from __future__ import annotations

import io
import re
import time
import uuid
from dataclasses import dataclass, field
from typing import Optional

import config

_MAX_DOC_BYTES = 20 * 1024 * 1024

# llama-text-embed-v2 truncates at 2048 tokens, so the ceiling is not the
# binding constraint -- precision is. Business documents pack several unrelated
# facts per page, and a chunk spanning all of them embeds to their average,
# which matches every query weakly and none strongly. Small chunks keep one
# topic per vector; the overlap stops a fact dying on a boundary.
_CHUNK_CHARS = 700
_CHUNK_OVERLAP = 120

SUFFIXES = (".pdf", ".docx", ".txt", ".md", ".xlsx", ".xlsm", ".xls", ".csv",
            ".pptx")


class RagError(RuntimeError):
    """Raised for anything the caller should surface to the user."""


@dataclass
class Document:
    id: str
    company_id: str
    filename: str
    chunks: int
    chars: int
    uploaded_at: float = field(default_factory=time.time)
    # The chunk text is kept so a document can be re-indexed under another
    # company without the user uploading the file again. Pinecone can be read
    # back, but only via list+fetch per id; holding the text costs a few KB and
    # makes carrying context forward a plain re-upsert.
    texts: list = field(default_factory=list)


# Documents are listed from this registry rather than from Pinecone, which is
# eventually consistent: a file uploaded a second ago may not yet appear in a
# query, and a UI that showed it vanishing would look broken.
_DOCS: dict[str, Document] = {}


# --- extraction ----------------------------------------------------------------

def _pdf_text(content: bytes) -> str:
    import pdfplumber
    out = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            out.append(page.extract_text() or "")
    return "\n".join(out)


def _docx_text(content: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(content))
    parts = [p.text for p in d.paragraphs]
    # Tables carry the headcount and site lists that matter most here, so they
    # are flattened into rows rather than skipped.
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pptx_text(content: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(content))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"[slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
    return "\n".join(parts)


def _sheet_text(content: bytes, filename: str) -> str:
    import pandas as pd
    if filename.lower().endswith(".csv"):
        frames = {"csv": pd.read_csv(io.BytesIO(content), dtype=str)}
    else:
        frames = pd.read_excel(io.BytesIO(content), sheet_name=None, dtype=str)
    parts = []
    for name, df in frames.items():
        df = df.fillna("")
        parts.append(f"[sheet {name}]")
        parts.append(" | ".join(str(c) for c in df.columns))
        for _, row in df.iterrows():
            cells = [str(v).strip() for v in row.tolist()]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(content: bytes, filename: str) -> str:
    """Plain text from any supported document."""
    name = (filename or "").lower()
    if not content:
        raise RagError("The uploaded file is empty.")
    if len(content) > _MAX_DOC_BYTES:
        raise RagError(
            f"File too large: {len(content) / 1e6:.1f} MB. The limit is "
            f"{_MAX_DOC_BYTES // (1024 * 1024)} MB.")
    try:
        if name.endswith(".pdf"):
            text = _pdf_text(content)
        elif name.endswith(".docx"):
            text = _docx_text(content)
        elif name.endswith(".pptx"):
            text = _pptx_text(content)
        elif name.endswith((".xlsx", ".xlsm", ".xls", ".csv")):
            text = _sheet_text(content, name)
        elif name.endswith((".txt", ".md")):
            text = content.decode("utf-8", errors="replace")
        else:
            raise RagError(
                f"Unsupported file type. Accepted: {', '.join(SUFFIXES)}.")
    except RagError:
        raise
    except Exception as e:
        raise RagError(f"Could not read {filename}: {e}") from e

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise RagError(
            f"No readable text in {filename}. A scanned PDF needs OCR first.")
    return text


# --- chunking ------------------------------------------------------------------

def chunk(text: str) -> list[str]:
    """Split on paragraph boundaries, packing up to the chunk budget.

    Paragraphs are kept whole where they fit so a chunk reads as prose rather
    than a fragment; only a single oversized paragraph is cut mid-way.
    """
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    buf = ""
    for para in paragraphs:
        while len(para) > _CHUNK_CHARS:
            if buf:
                chunks.append(buf)
                buf = ""
            chunks.append(para[:_CHUNK_CHARS])
            para = para[_CHUNK_CHARS - _CHUNK_OVERLAP:]
        if not buf:
            buf = para
        elif len(buf) + len(para) + 2 <= _CHUNK_CHARS:
            buf = f"{buf}\n\n{para}"
        else:
            chunks.append(buf)
            buf = para
    if buf:
        chunks.append(buf)
    return chunks


# --- pinecone ------------------------------------------------------------------

_index = None


def _get_index():
    global _index
    if _index is not None:
        return _index
    if not config.PINECONE_API_KEY:
        raise RagError(
            "Pinecone is not configured: set PINECONE_APIKEY in backend/.env.")
    try:
        from pinecone import Pinecone
    except ImportError as e:
        raise RagError("The `pinecone` package is not installed.") from e
    pc = Pinecone(api_key=config.PINECONE_API_KEY)
    name = config.PINECONE_INDEX
    if not pc.has_index(name):
        raise RagError(
            f"Pinecone index '{name}' does not exist. Create it with an "
            f"integrated embedding model, or set PINECONE_INDEX.")
    _index = pc.Index(name)
    return _index


def configured() -> bool:
    return bool(config.PINECONE_API_KEY)


def index_document(company_id: str, content: bytes, filename: str) -> Document:
    """Extract, chunk and upsert one document into the company's namespace."""
    text = extract_text(content, filename)
    chunks = chunk(text)
    if not chunks:
        raise RagError(f"No readable text in {filename}.")

    idx = _get_index()
    doc_id = uuid.uuid4().hex[:12]
    records = [
        {
            "_id": f"{doc_id}-{i}",
            "text": c,
            "doc_id": doc_id,
            "filename": filename,
            "chunk": i,
        }
        for i, c in enumerate(chunks)
    ]
    # Pinecone caps a single upsert batch, so long documents go up in slices.
    for start in range(0, len(records), 90):
        idx.upsert_records(namespace=company_id,
                           records=records[start:start + 90])

    doc = Document(id=doc_id, company_id=company_id, filename=filename,
                   chunks=len(chunks), chars=len(text), texts=chunks)
    _DOCS[doc_id] = doc
    return doc


def documents(company_id: str) -> list[Document]:
    return sorted((d for d in _DOCS.values() if d.company_id == company_id),
                  key=lambda d: d.uploaded_at)


def delete_document(company_id: str, doc_id: str) -> bool:
    doc = _DOCS.get(doc_id)
    if not doc or doc.company_id != company_id:
        return False
    idx = _get_index()
    idx.delete(namespace=company_id,
               ids=[f"{doc_id}-{i}" for i in range(doc.chunks)])
    _DOCS.pop(doc_id, None)
    return True


def search(company_id: str, query: str, top_k: int = 5) -> list[dict]:
    """Top matching chunks from this company's namespace only."""
    if not query.strip():
        return []
    idx = _get_index()
    try:
        res = idx.search(namespace=company_id,
                         query={"inputs": {"text": query}, "top_k": top_k})
    except Exception as e:
        raise RagError(f"Pinecone search failed: {e}") from e
    hits = []
    for h in res["result"]["hits"]:
        fields = h["fields"]
        hits.append({
            "score": h["score"],
            "text": fields.get("text", ""),
            "filename": fields.get("filename", "?"),
            "chunk": fields.get("chunk"),
        })
    return hits


def all_chunks(company_id: str) -> list[dict]:
    """Every chunk in the company's namespace, in document order.

    Indexing the same file twice leaves two copies of each chunk; they are
    deduplicated on (filename, chunk) so the prompt never carries a pack twice.
    """
    idx = _get_index()
    ids: list[str] = []
    try:
        for batch in idx.list(namespace=company_id):
            for item in batch:
                ids.append(item.id if hasattr(item, "id") else str(item))
    except Exception as e:
        raise RagError(f"Pinecone list failed: {e}") from e
    if not ids:
        return []
    seen: dict = {}
    for start in range(0, len(ids), 200):
        try:
            res = idx.fetch(ids=ids[start:start + 200], namespace=company_id)
        except Exception as e:
            raise RagError(f"Pinecone fetch failed: {e}") from e
        vecs = res.vectors if hasattr(res, "vectors") else res["vectors"]
        for v in vecs.values():
            md = v.metadata if hasattr(v, "metadata") else v.get("metadata")
            if not md or not md.get("text"):
                continue
            key = (str(md.get("filename", "?")), int(md.get("chunk") or 0))
            seen.setdefault(key, {"filename": key[0], "chunk": key[1],
                                  "text": md["text"]})
    return [seen[k] for k in sorted(seen)]


_HEADER = (
    "COMPANY DOCUMENTS - the company's own records. Facts here (sites, "
    "occupancy, headcount, systems, leases, contracts, duplicated teams) are "
    "established: use them, name the file they came from, and do not ask the "
    "user for them. Before saying you do not have a figure, check every file "
    "below - a contract value, a system cost or a headcount is usually in a "
    "different file from the one being discussed. Dollar amounts here are "
    "costs or contract values, never savings - do not turn them into an "
    "estimate of what could be saved.")


def render_context(company_id: str, query: str, top_k: int = 5,
                   min_score: float = 0.15) -> Optional[str]:
    """The document block as it appears in the agent prompt.

    A small pack goes in whole. Retrieval returns the chunks nearest the
    current topic, so while a conversation was on facilities the vendor
    register and the IT-systems file were never in the prompt -- and the
    agent, correctly for what it could see, said "there's no contract value
    in what I have" about a contract the pack listed. Up to
    RAG_FULL_CONTEXT_MAX chunks, the whole pack costs a few thousand tokens
    and removes that failure outright. Above it, the query's best hits are
    joined by the opening chunk of every file not already represented, so no
    document is invisible on a given turn.
    """
    if not configured():
        return None
    try:
        chunks = all_chunks(company_id)
    except RagError:
        chunks = []
    if not chunks:
        return None

    if len(chunks) <= config.RAG_FULL_CONTEXT_MAX:
        lines = [_HEADER, ""]
        current = None
        for c in chunks:
            if c["filename"] != current:
                current = c["filename"]
                lines.append(f"## {current}")
            lines.append(c["text"])
        return "\n".join(lines)

    try:
        hits = [h for h in search(company_id, query, top_k)
                if h["score"] >= min_score]
    except RagError:
        hits = []
    chosen = {(h["filename"], int(h["chunk"] or 0)): h["text"] for h in hits}
    covered = {h["filename"] for h in hits}
    for c in chunks:
        if c["filename"] not in covered and c["chunk"] == 0:
            chosen[(c["filename"], 0)] = c["text"]
    if not chosen:
        return None
    lines = [_HEADER, ""]
    for (fname, _), text in sorted(chosen.items()):
        lines.append(f"- [{fname}] {text}")
    return "\n".join(lines)


def copy_documents(source_company_id: str, target_company_id: str) -> int:
    """Re-index one company's documents under another.

    Uploading a P&L creates a NEW company, whose document namespace starts
    empty -- so context gathered against the previous company silently stopped
    reaching the agent, and it went back to asking questions the documents
    already answered. Carrying the documents across keeps the namespace
    boundary intact (nothing is shared at query time) while matching what a
    single operator diagnosing one business actually expects.
    """
    if source_company_id == target_company_id:
        return 0
    src = documents(source_company_id)
    if not src:
        return 0
    idx = _get_index()
    copied = 0
    for doc in src:
        if not doc.texts:
            continue
        new_id = uuid.uuid4().hex[:12]
        records = [
            {"_id": f"{new_id}-{i}", "text": c, "doc_id": new_id,
             "filename": doc.filename, "chunk": i}
            for i, c in enumerate(doc.texts)
        ]
        for start in range(0, len(records), 90):
            idx.upsert_records(namespace=target_company_id,
                               records=records[start:start + 90])
        _DOCS[new_id] = Document(
            id=new_id, company_id=target_company_id, filename=doc.filename,
            chunks=len(doc.texts), chars=doc.chars, texts=doc.texts)
        copied += 1
    return copied
