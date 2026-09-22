"""Candidate tables from documents that are not spreadsheets.

PDFs and decks are where finance teams actually keep a P&L when they are not
sending a spreadsheet. This module finds the tables in them and returns RAW
CELL GRIDS -- strings, exactly as printed. It never converts a figure, picks a
line item or decides what anything means: that stays in ingest.py, so a number
pulled off a slide goes through the same mapping-confirmation as a CSV cell
before it reaches a benchmark.

What is deliberately NOT handled: scanned pages and screenshots. Recognising
digits from pixels is a different risk class -- a misread 8 for a 3 becomes a
confident wrong benchmark -- so those are refused with instructions instead.
"""

import io
import re
from typing import Callable, List, Tuple

Grid = List[List[str]]
# (section label, cell grid, surrounding text). The text carries what the table
# itself does not: "(USD in thousands)" is printed above a statement, never
# inside it, and reading it as dollars understates the company 1000-fold.
Candidate = Tuple[str, Grid, str]

# A P&L fragment worth offering: a label column, at least one figure column,
# and enough rows to be a statement rather than a stray two-cell box.
MIN_ROWS = 3
MIN_COLS = 2

_YEARISH = re.compile(r"(?:19|20)\d{2}")
# Whitespace-aligned columns: "General & administrative      44,950   41,500".
_COLUMN_GAP = re.compile(r"\s{2,}|\t")


def _clean(grid) -> Grid:
    """Strings, stripped, with fully empty rows and columns removed."""
    rows = [[("" if c is None else str(c)).replace("\n", " ").strip()
             for c in (row or [])] for row in (grid or [])]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [i for i in range(width) if any(r[i] for r in rows)]
    return [[r[i] for i in keep] for r in rows]


def _usable(grid: Grid) -> bool:
    return len(grid) >= MIN_ROWS and bool(grid) and len(grid[0]) >= MIN_COLS


def _label(prefix: str, index: int, total: int) -> str:
    return prefix if total == 1 else f"{prefix}, table {index}"


def _aligned_columns(text: str, is_number: Callable) -> Grid:
    """Rebuild a table from a page that has no ruled lines.

    Most statement PDFs draw no cell borders, so pdfplumber finds no table at
    all -- and the extracted text separates columns with a single space, so
    splitting on whitespace runs does not work either ("Net product revenue
    310,000 340,000"). What is reliable: figures sit at the END of the line. So
    peel numeric tokens off the right; whatever remains is the line item. The
    period header is the line whose trailing tokens are years.
    """
    rows: Grid = []
    header: List[str] = []
    for raw in text.splitlines():
        tokens = raw.strip().split()
        if len(tokens) < 2:
            continue

        values: List[str] = []
        while tokens and is_number(tokens[-1]) is not None:
            values.insert(0, tokens.pop())
        label = " ".join(tokens).strip()

        if values and label:
            rows.append([label, *values])
            continue

        if not rows and not values:
            years: List[str] = []
            while tokens and _YEARISH.search(tokens[-1]):
                years.insert(0, tokens.pop())
            if years:
                header = [" ".join(tokens).strip(), *years]

    if not rows:
        return []

    # Keep the shape the statement actually has: a stray footer such as
    # "Page 1 of 3" leaves one value where every real line has two.
    counts = [len(r) for r in rows]
    width = max(set(counts), key=counts.count)
    rows = [r for r in rows if len(r) == width]
    if header and len(header) == width:
        rows.insert(0, header)
    return _clean(rows)


def pdf_tables(content: bytes, is_number: Callable) -> List[Candidate]:
    """Every table-like block in a PDF, page by page."""
    import pdfplumber

    out: List[Candidate] = []
    any_text = False
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        if not pdf.pages:
            raise ValueError("That PDF has no pages.")
        for number, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            any_text = any_text or bool(text.strip())
            found = [g for g in (_clean(t) for t in (page.extract_tables() or []))
                     if _usable(g)]
            for i, grid in enumerate(found, 1):
                out.append((_label(f"Page {number}", i, len(found)), grid, text))
            if not found and text.strip():
                grid = _aligned_columns(text, is_number)
                if _usable(grid):
                    out.append((f"Page {number}", grid, text))

    if out:
        return out
    if not any_text:
        raise ValueError(
            "This PDF has no text in it -- it is a scan or a photo of a "
            "document. Reading digits from an image is not supported, because a "
            "misread figure would become a confident wrong benchmark. Export "
            "the P&L from the system that produced it, as CSV or Excel.")
    raise ValueError(
        "No table of figures was found in that PDF. If the P&L is there, "
        "export it as CSV or Excel instead.")


def pptx_tables(content: bytes, is_number: Callable = None) -> List[Candidate]:
    """Every native table in a deck, slide by slide.

    Tables pasted as pictures are invisible here, as they should be.
    """
    from pptx import Presentation

    prs = Presentation(io.BytesIO(content))
    out: List[Candidate] = []
    for number, slide in enumerate(prs.slides, 1):
        found, words = [], []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                grid = _clean([[cell.text for cell in row.cells]
                               for row in shape.table.rows])
                if _usable(grid):
                    found.append(grid)
            elif getattr(shape, "has_text_frame", False):
                words.append(shape.text_frame.text)
        text = " ".join(w for w in words if w)
        for i, grid in enumerate(found, 1):
            out.append((_label(f"Slide {number}", i, len(found)), grid, text))

    if out:
        return out
    raise ValueError(
        "No table was found in that deck. Figures sitting in text boxes or "
        "pasted in as images cannot be read reliably: export the P&L as CSV "
        "or Excel instead.")
